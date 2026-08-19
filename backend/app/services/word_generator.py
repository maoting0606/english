from io import BytesIO
from docx import Document
from app.services.parser import WordItem

def build_docx(items: list[WordItem], dictation: bool = False) -> bytes:
    doc = Document()
    doc.add_heading("单词默写" if dictation else "单词原版对照", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "英文"
    table.rows[0].cells[1].text = "中文"
    for item in items:
        row = table.add_row().cells
        row[0].text = "__________" if dictation else item.word
        row[1].text = item.meaning
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
